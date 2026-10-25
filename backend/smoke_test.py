"""End-to-end smoke test for the interview lifecycle (runs without a key).

    .venv/bin/python smoke_test.py
"""
from __future__ import annotations

import io
import json
import sys
import time

from fastapi.testclient import TestClient

from app.main import app

FAILURES = []


def check(label: str, cond: bool, extra: str = "") -> None:
    print(("  PASS  " if cond else "  FAIL  ") + label + (f"  {extra}" if extra else ""))
    if not cond:
        FAILURES.append(label)


def make_resume_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Ananya Sharma", 0)
    doc.add_paragraph("ananya.sharma@example.com | +91 98765 43210")
    doc.add_paragraph("Software Engineer with 4 years of experience building "
                      "distributed services in Java and Python.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Java, Spring Boot, Microservices, Docker, Kubernetes, "
                      "PostgreSQL, REST APIs, JUnit, Kafka")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Software Engineer, Nexcore Systems, 2021 - Present")
    doc.add_paragraph("- Built a Spring Boot order service handling 3k rps")
    doc.add_paragraph("- Migrated a monolith to 6 microservices on Kubernetes")
    doc.add_heading("Projects", level=1)
    doc.add_paragraph("Inventory Sync Pipeline - Kafka based reconciliation tool")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.Tech Computer Science, Shiv Nadar University, 2021")
    doc.add_heading("Certifications", level=1)
    doc.add_paragraph("AWS Certified Solutions Architect - Associate")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


JD = """Senior Backend Engineer

We are hiring a backend engineer to own our payments platform. You will design
and operate high-throughput REST and event-driven services, improve reliability,
and mentor junior engineers.

Requirements:
- 4+ years building production services in Java or Python
- Strong with Spring Boot, REST API design and microservices
- Experience with PostgreSQL, Kafka, Redis and caching strategies
- Docker and Kubernetes in production
- Testing culture: unit and integration tests, CI/CD
- Good written and verbal communication
"""


def main() -> int:
    from app.db import init_db

    init_db()
    client = TestClient(app)

    print("\n== health ==")
    r = client.get("/api/health")
    check("health 200", r.status_code == 200, str(r.status_code))
    health = r.json()
    check("llm status present", "llm" in health, json.dumps(health.get("llm")))

    print("\n== create interview (JD + DOCX resume) ==")
    r = client.post(
        "/api/interviews",
        data={
            "job_title": "Senior Backend Engineer",
            "job_description": JD,
            "required_skills": "Java, Spring Boot, Kafka, Kubernetes",
            "experience_level": "senior",
            "question_count": "10",
            "categories": "technical,behavioral,experience,scenario",
            "difficulty": "mixed",
            "duration_minutes": "30",
            "per_question_limit_seconds": "90",
            "allow_followups": "true",
            "max_followups": "2",
            "weight_technical": "60",
            "weight_behavioral": "10",
            "weight_integrity": "30",
            "recruiter_questions": "1. What is your experience with payment idempotency?",
        },
        files={"resume": ("ananya_sharma.docx", make_resume_docx(),
                          "application/vnd.openxmlformats-officedocument."
                          "wordprocessingml.document")},
    )
    check("create 200", r.status_code == 200, f"{r.status_code} {r.text[:300]}")
    if r.status_code != 200:
        return 1
    iv = r.json()

    check("resume parsed", bool(iv["resume"]), iv["resume"]["candidate_name"])
    check("candidate name extracted", iv["resume"]["candidate_name"] == "Ananya Sharma",
          repr(iv["resume"]["candidate_name"]))
    check("resume skills extracted",
          "java" in [s.lower() for s in iv["resume"]["parsed"]["skills"]],
          str(iv["resume"]["parsed"]["skills"][:8]))
    scored = lambda qs: [q for q in qs if q["category"] != "icebreaker"]
    check("default question count = 10 (ice-breakers excluded)",
          len(scored(iv["questions"])) == 10, f"got {len(scored(iv['questions']))}")
    check("1-2 unscored ice-breakers precede the core set",
          1 <= len(iv["questions"]) - len(scored(iv["questions"])) <= 2)
    check("recruiter weights stored",
          iv.get("weights") == {"technical": 60, "behavioral": 10, "integrity": 30},
          str(iv.get("weights")))
    recruiter_q = [q for q in iv["questions"] if q["source"] == "recruiter"]
    check("recruiter question kept", len(recruiter_q) == 1, str(len(recruiter_q)))
    texts = [q["text"].lower() for q in iv["questions"]]
    check("no duplicate questions", len(set(texts)) == len(texts))
    cats = {q["category"] for q in iv["questions"]}
    check("multiple categories present", len(cats) >= 3, str(sorted(cats)))
    skills_targeted = {q["skill_target"] for q in iv["questions"]}
    check("questions target real skills",
          any("kafka" in s.lower() or "spring" in s.lower() or "java" in s.lower()
              for s in skills_targeted), str(list(skills_targeted)[:5]))

    print("\n== regenerate with a different count ==")
    r = client.post(f"/api/interviews/{iv['id']}/regenerate",
                    json={"question_count": 5})
    check("regenerate 200", r.status_code == 200, str(r.status_code))
    check("5 scored questions after regenerate",
          len(scored(r.json()["questions"])) == 5,
          f"got {len(scored(r.json()['questions']))}")

    r = client.post(f"/api/interviews/{iv['id']}/regenerate",
                    json={"question_count": 10})
    check("restored to 10", len(scored(r.json()["questions"])) == 10,
          f"got {len(scored(r.json()['questions']))}")

    print("\n== session: onboarding -> interview ==")
    r = client.post("/api/sessions", json={"interview_id": iv["id"],
                                           "candidate_name": ""})
    check("session created", r.status_code == 200, r.text[:200])
    s = r.json()
    check("candidate name from resume", s["candidate_name"] == "Ananya Sharma",
          s["candidate_name"])
    sid = s["id"]

    r = client.post(f"/api/sessions/{sid}/consent",
                    json={"consent_recording": True, "consent_proctoring": True,
                          "device_check": {"camera": "ok", "microphone": "ok",
                                           "resolution": "1280x720"}})
    check("consent stored", r.status_code == 200 and
          r.json()["consent_proctoring"] is True)

    r = client.get(f"/api/sessions/{sid}/pair")
    check("pair token issued", r.status_code == 200 and len(r.json()["token"]) > 20)
    token = r.json()["token"]
    r = client.get(f"/api/pair/{token}")
    check("pair token resolves", r.status_code == 200, r.text[:200])

    r = client.post(f"/api/sessions/{sid}/start")
    check("session started", r.status_code == 200, r.text[:200])
    sess = r.json()
    check("answer rows created (10 scored + ice-breakers)",
          len(scored(sess["answers"])) == 10,
          f"got {len(scored(sess['answers']))}")

    print("\n== answers + follow-up ==")
    answers = sess["answers"]
    r = client.post(f"/api/sessions/{sid}/answers/{answers[0]['id']}",
                    json={"transcript": "In my last role I owned a Spring Boot "
                          "order service that handled 3000 requests per second. "
                          "We used Kafka to decouple inventory updates and "
                          "PostgreSQL for persistence, with JUnit integration "
                          "tests running in CI on every merge.",
                          "status": "answered", "duration_seconds": 64})
    check("answer saved", r.status_code == 200 and
          r.json()["status"] == "answered", r.text[:200])

    r = client.post(f"/api/sessions/{sid}/answers/{answers[1]['id']}",
                    json={"transcript": "", "status": "skipped"})
    check("skip recorded", r.json()["status"] == "skipped")

    r = client.post(f"/api/sessions/{sid}/followup",
                    json={"parent_answer_id": answers[0]["id"]})
    check("follow-up handled", r.status_code == 200, r.text[:200])
    fu = r.json()
    check("follow-up allowed or reason given",
          "allowed" in fu, f"allowed={fu.get('allowed')} {fu.get('reason')}")

    for a in answers[2:]:
        client.post(f"/api/sessions/{sid}/answers/{a['id']}",
                    json={"transcript": "We would first measure the bottleneck "
                          "with traces, then isolate the service, roll back the "
                          "bad deploy and add a regression test before "
                          "redeploying behind a feature flag.",
                          "status": "answered", "duration_seconds": 40})

    r = client.get(f"/api/sessions/{sid}/next")
    prog = r.json()
    check("progress tracking", prog["answered"] >= 9,
          json.dumps({k: prog[k] for k in ("total", "answered", "skipped", "unanswered")}))

    print("\n== proctoring incidents ==")
    frame = ("data:image/jpeg;base64,"
             "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8"
             "z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
    events = [
        {"type": "face_absent", "ts_ms": 45000, "value": "4.2s",
         "duration_ms": 4200, "evidence_data_url": frame},
        {"type": "multiple_faces", "ts_ms": 91000, "value": "2 faces",
         "duration_ms": 3100, "evidence_data_url": frame},
        {"type": "secondary_voice", "ts_ms": 122000, "value": "speech 0.71"},
        {"type": "gaze_aversion", "ts_ms": 150000, "value": "yaw 31deg"},
        {"type": "camera_lost", "ts_ms": 200000},
    ]
    for ev in events:
        r = client.post(f"/api/sessions/{sid}/incidents", json=ev)
        check(f"log {ev['type']}", r.status_code == 200, r.text[:200])

    r = client.get(f"/api/sessions/{sid}/incidents", params={"evidence": True})
    inc = r.json()
    check("5 incidents stored", len(inc) == 5, str(len(inc)))
    check("observation separated from inference",
          all(i["observed"] and i["inference"] for i in inc),
          inc[0]["observed"][:60] + " | " + inc[0]["inference"][:60])
    check("evidence attached", inc[0].get("evidence_data_url", "").startswith("data:image"))
    check("severity assigned", all(i["severity"] in ("info", "warning", "high")
                                   for i in inc), str([i["severity"] for i in inc]))

    print("\n== backup question (camera glitch) ==")
    r = client.post(f"/api/sessions/{sid}/backup")
    b = r.json()
    check("backup deployed", r.status_code == 200 and b.get("allowed") is True,
          r.text[:200])
    check("backup is a pending scored question",
          b["answer"]["status"] == "pending" and
          b["answer"]["category"] != "icebreaker" and
          b["answer"]["question_text"].startswith("Backup round:"))
    r = client.post(f"/api/sessions/{sid}/backup")
    check("second backup allowed", r.json().get("allowed") is True)
    r = client.post(f"/api/sessions/{sid}/backup")
    check("backup budget capped at 2", r.json().get("allowed") is False and
          r.json().get("reason") == "backup_budget_exhausted", r.text[:200])

    print("\n== complete + evaluation ==")
    r = client.post(f"/api/sessions/{sid}/complete")
    check("completed", r.status_code == 200 and r.json()["status"] == "completed")

    deadline = time.time() + 60
    report = None
    while time.time() < deadline:
        report = client.get(f"/api/sessions/{sid}/report").json()
        if report["evaluation_status"] in ("done", "failed"):
            break
        time.sleep(0.5)
    check("evaluation finished", report["evaluation_status"] == "done",
          report["evaluation_status"])
    ev = report["evaluation"] or {}
    check("per-question scores", len(ev.get("per_question", [])) >= 10,
          f"got {len(ev.get('per_question', []))}")
    scores = [p["score"] for p in ev.get("per_question", [])]
    check("scores within 1-5", all(1 <= s <= 5 for s in scores), str(scores))
    check("evidence quotes present",
          sum(1 for p in ev.get("per_question", []) if p.get("evidence")) >= 5)
    check("dimension scores", set(ev.get("dimension_scores", {})) >=
          {"technical_knowledge", "problem_solving", "communication", "role_fit"},
          str(list(ev.get("dimension_scores", {}))))
    check("summary written", len(ev.get("summary", "")) > 50)
    check("proctoring summary present", report["proctoring_summary"]["total"] == 5)
    check("no-verdict disclaimer", "not proof" in
          report["proctoring_summary"]["disclaimer"].lower())
    check("answers include transcripts",
          all("transcript" in a for a in report["answers"]))

    print("\n== scorecard (Technical /10, Behavioral /5, Integrity index) ==")
    card = report.get("scorecard") or {}
    check("scorecard present", bool(card))
    check("technical score is /10",
          card.get("technical_10") is not None and 0 <= card["technical_10"] <= 10,
          str(card.get("technical_10")))
    check("behavioral score is /5",
          card.get("behavioral_5") is not None and 0 <= card["behavioral_5"] <= 5,
          str(card.get("behavioral_5")))
    # 1 camera_lost (high 8) + 2 high (16) ... compute from stored incidents
    expected = 100 - sum(8 if i["severity"] == "high" else 3 if i["severity"] == "warning" else 0
                         for i in report["incidents"])
    check("integrity index = 100 - weighted incidents",
          card.get("integrity_index") == expected,
          f"{card.get('integrity_index')} vs {expected}")
    check("recruiter weights applied", card["weights"]["technical"] == 60)
    check("ice-breakers excluded from scoring",
          card["technical_questions"] + card["behavioral_questions"] ==
          len([a for a in report["answers"] if a["category"] != "icebreaker"
               and a["status"] != "pending"]), str(card["technical_questions"]))
    check("veto carries a reason", bool(card["veto"]["reason"]))

    print("\n== recruiter override ==")
    target = report["answers"][0]["id"]
    r = client.patch(f"/api/sessions/{sid}",
                     json={"recruiter_notes": "Strong backend depth.",
                           "overrides": [{"answer_id": target, "score": 5,
                                          "note": "Verified in round 2"}],
                           "acknowledged_incidents": [inc[0]["id"]]})
    check("override saved", r.status_code == 200, r.text[:300])
    out = r.json()
    check("override visible in report",
          any(a["id"] == target and a.get("score") == 5 for a in out["answers"]))
    check("incident acknowledged",
          any(i["id"] == inc[0]["id"] and i["acknowledged"] for i in out["incidents"]))

    r = client.patch(f"/api/sessions/{sid}", json={"veto_override": "reinstated"})
    check("veto override saved", r.status_code == 200 and
          r.json()["scorecard"]["veto"]["final"] == "reinstated", r.text[:200])
    r = client.patch(f"/api/sessions/{sid}", json={"veto_override": "bogus"})
    check("bad veto value rejected", r.status_code == 400)

    print("\n== recruitment dashboard listing ==")
    r = client.get("/api/interviews")
    check("interviews listed", r.status_code == 200 and len(r.json()) >= 1)
    r = client.get(f"/api/sessions/{sid}")
    check("session fetchable", r.status_code == 200 and
          r.json()["incident_count"] == 5)

    print()
    if FAILURES:
        print(f"{len(FAILURES)} FAILURE(S):")
        for f in FAILURES:
            print("  - " + f)
        return 1
    print("ALL CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
