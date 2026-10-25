#!/usr/bin/env python3
"""Generate the demo assets used for an end-to-end walk-through.

Run from anywhere:

    python3 scripts/make_samples.py

Produces:
    samples/sample_job_description.txt
    samples/sample_resume.docx      (python-docx)
    samples/sample_resume.pdf       (hand-written minimal PDF, pypdf-readable)
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"

JOB_TITLE = "Software Engineer — Backend (Java / Spring Boot)"

JOB_DESCRIPTION = """Software Engineer — Backend (Java / Spring Boot)
Location: Bengaluru / Hybrid · Experience: 2–4 years · Full time

About the role
We are hiring a backend engineer to build and operate the services that power
our payments platform. You will own microservices end to end: design, implement,
test, deploy and monitor them in production.

Responsibilities
- Design and implement REST and event-driven microservices in Java 17+ and Spring Boot.
- Model data in PostgreSQL and write efficient, well-tuned queries.
- Build resilient integrations with Kafka, Redis and third-party payment gateways.
- Improve latency, throughput and observability of existing services.
- Participate in code reviews, incident response and capacity planning.
- Work with product and frontend engineers to deliver features incrementally.

Required skills
- Java, Spring Boot, Spring Security
- REST API design and microservices architecture
- SQL / PostgreSQL, at least one message broker (Kafka or RabbitMQ)
- Git, unit and integration testing (JUnit, Mockito)
- Docker and a CI/CD pipeline

Nice to have
- Kubernetes, AWS (ECS/Lambda), Redis
- Event sourcing or CQRS exposure
- Experience with PCI-DSS or other regulated environments

What we evaluate in interviews
- Depth of Java and Spring Boot knowledge (collections, concurrency, dependency injection)
- Microservice design: boundaries, failure modes, idempotency, retries
- Practical experience from past projects
- Debugging and problem-solving approach
- Communication clarity and behavioural consistency
"""

RESUME_LINES = [
    ("H1", "Ananya Sharma"),
    ("SUB", "ananya.sharma@example.com · +91 98765 43210 · Bengaluru, India · github.com/ananyasharma"),
    ("H2", "Summary"),
    ("P", "Backend software engineer with 3 years of experience building Java and "
          "Spring Boot microservices for payments and e-commerce domains. Worked on "
          "REST APIs, Kafka-based event pipelines and PostgreSQL performance tuning."),
    ("H2", "Skills"),
    ("P", "Java 17, Spring Boot, Spring Security, Spring Data JPA, REST, microservices, "
          "PostgreSQL, MySQL, Kafka, RabbitMQ, Redis, Docker, Git, JUnit, Mockito, "
          "Jenkins CI, AWS (ECS, S3)"),
    ("H2", "Experience"),
    ("H3", "Software Engineer II — Payverse Technologies (2023 – Present)"),
    ("P", "- Built a refund processing service in Spring Boot handling 40k refunds/day, "
          "cutting settlement latency by 35%."),
    ("P", "- Designed idempotent payment webhook consumers on Kafka with retry and "
          "dead-letter handling."),
    ("P", "- Migrated 12 REST endpoints to a versioned API contract and added Mockito "
          "integration tests raising coverage from 54% to 82%."),
    ("H3", "Software Engineer — Shopkart Solutions (2021 – 2023)"),
    ("P", "- Developed order and inventory microservices in Java/Spring Boot behind an "
          "API gateway."),
    ("P", "- Optimised slow PostgreSQL queries (missing indexes, N+1 ORM fetches) and "
          "reduced p95 API latency from 900 ms to 320 ms."),
    ("P", "- Containerised services with Docker and maintained the Jenkins pipeline."),
    ("H2", "Projects"),
    ("P", "RideShare Analytics — Spring Boot + Kafka + Redis service that aggregates "
          "ride events in near real time; deployed on AWS ECS."),
    ("P", "FileVault — Spring Security based document vault with JWT refresh tokens "
          "and per-object access policies."),
    ("H2", "Education"),
    ("P", "B.E. Computer Science, RV College of Engineering, 2021 — CGPA 8.4/10"),
    ("H2", "Certifications"),
    ("P", "AWS Certified Cloud Practitioner (2023) · Oracle Certified Professional: "
          "Java SE 11 Developer (2022)"),
]


def write_job_description() -> Path:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    path = SAMPLES / "sample_job_description.txt"
    path.write_text(JOB_DESCRIPTION.strip() + "\n", encoding="utf-8")
    return path


def write_docx() -> Path:
    from docx import Document  # lazy: only needed when generating

    doc = Document()
    for kind, text in RESUME_LINES:
        if kind == "H1":
            doc.add_heading(text, level=0)
        elif kind == "H2":
            doc.add_heading(text, level=1)
        elif kind == "H3":
            doc.add_heading(text, level=2)
        elif kind == "SUB":
            para = doc.add_paragraph(text)
            para.runs[0].italic = True
        else:
            doc.add_paragraph(text)
    path = SAMPLES / "sample_resume.docx"
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return path


def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def write_pdf() -> Path:
    """Write a minimal single-page PDF (Helvetica) that pypdf can extract."""
    width, height = 595, 842  # A4 in points
    y = height - 56
    leading_map = {"H1": 26, "H2": 24, "H3": 19, "SUB": 18, "P": 16}

    ops = ["BT"]
    for kind, text in RESUME_LINES:
        size = {"H1": 20, "H2": 13, "H3": 11.5, "SUB": 9.5, "P": 10}[kind]
        font = "F2" if kind in ("H1", "H2", "H3") else "F1"
        # crude wrap so nothing runs off the page
        max_chars = 92
        chunks = [text[i : i + max_chars] for i in range(0, len(text), max_chars)] or [""]
        for i, chunk in enumerate(chunks):
            if i:
                y -= leading_map[kind]
            ops.append(
                f"/{font} {size} Tf 1 0 0 1 56 {y:.1f} Tm ({_pdf_escape(chunk)}) Tj"
            )
        y -= leading_map[kind]
        if y < 70:
            break
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1", "replace")

    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width} {height}] "
            f"/Resources << /Font << /F1 5 0 R /F2 6 0 R >> >> "
            f"/Contents 4 0 R >>"
        ).encode("latin-1"),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>",
    ]

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode())
        out.write(body)
        out.write(b"\nendobj\n")

    xref_at = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n".encode()
    )

    path = SAMPLES / "sample_resume.pdf"
    path.write_bytes(out.getvalue())

    # sanity check: the PDF parser must be able to read it back
    from pypdf import PdfReader

    text = "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    if "Ananya Sharma" not in text:
        raise SystemExit("generated PDF failed the pypdf round-trip check")
    return path


def main() -> int:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    made = [write_job_description(), write_docx(), write_pdf()]
    for p in made:
        print(f"✓ {p.relative_to(ROOT)} ({p.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
